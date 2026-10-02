#!/usr/bin/env python3
"""
Study Buddy
-----------
A local web app for your study material. Upload PDFs, Word files, slides, notes and notebooks;
they are stored on this computer, sorted by subject and kind automatically, summarized, and a
tutor agent teaches you from them. Everything runs offline and free through Ollama.

Usage:
  python app.py              # then open http://127.0.0.1:8001
  python app.py --port 8080

Files live in library/<id>/, metadata and search index in library.db.
Only listens on 127.0.0.1, so it's reachable from this computer only.
"""

import argparse
import hashlib
import json
import math
import os
import re
import sqlite3
import subprocess
import threading
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path

from flask import Flask, Response, abort, jsonify, request, send_from_directory, stream_with_context
from werkzeug.utils import secure_filename

HERE = Path(__file__).resolve().parent
LIB = HERE / "library"
WEB = HERE / "web"
DB_PATH = HERE / "library.db"
SETTINGS_PATH = HERE / "settings.json"
OLLAMA = "http://127.0.0.1:11434"

CLAUDE_MODEL = "claude-opus-5-5"   # optional, needs an Anthropic API key with credit (paid per use)
MODELS = ["qwen3:4b", "qwen3:8b", CLAUDE_MODEL]
EMBED_MODEL = "nomic-embed-text"   # small free embedding model for meaning-based search
FOUND_SIM = 0.58      # cosine similarity above which a passage is really about the question
KINDS = ["Notes", "Book / Textbook", "Slides", "Assignment", "Lab / Practical", "Question paper",
         "Code / Notebook", "Report", "Other"]
TEXT_EXTS = {".txt", ".md", ".csv", ".py", ".java", ".c", ".cpp", ".js", ".html", ".json", ".tex"}
EXTS = TEXT_EXTS | {".pdf", ".docx", ".pptx", ".ipynb"}

CHUNK = 1200          # characters per search chunk
MAX_PARTS = 16        # summary passes per document (long docs get bigger parts)
PART_MAX = 8000       # largest part, in characters, that still fits the context comfortably

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = 500 * 1024 ** 2

chat_count = {"n": 0}                   # live tutor chats; background work yields to them
chat_lock = threading.Lock()
work_queue = []
work_cv = threading.Condition()


# ---------------- settings & storage ----------------

def settings():
    s = {"model": MODELS[0]}
    if SETTINGS_PATH.exists():
        s.update(json.loads(SETTINGS_PATH.read_text()))
    return s


def db():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    LIB.mkdir(exist_ok=True)
    with db() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS materials (
            id TEXT PRIMARY KEY, filename TEXT, stored TEXT, sha TEXT, title TEXT, subject TEXT,
            kind TEXT, topics TEXT DEFAULT '[]', summary TEXT DEFAULT '', notes TEXT DEFAULT '[]',
            status TEXT, progress TEXT DEFAULT '', error TEXT DEFAULT '', chars INTEGER DEFAULT 0,
            created REAL);
        -- passages: search chunks that remember where they came from (page / slide) for citations
        CREATE VIRTUAL TABLE IF NOT EXISTS passages USING fts5(
            material_id UNINDEXED, idx UNINDEXED, page UNINDEXED, loc UNINDEXED, text);
        CREATE TABLE IF NOT EXISTS vectors (material_id TEXT, idx INTEGER, vec BLOB, PRIMARY KEY (material_id, idx));
        CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT, scope TEXT, role TEXT, content TEXT, created REAL);
        DROP TABLE IF EXISTS chunks;
        """)
        for table, col in (("materials", "unit"), ("materials", "semester"), ("messages", "sources")):
            if col not in [r[1] for r in c.execute(f"PRAGMA table_info({table})")]:
                c.execute(f"ALTER TABLE {table} ADD COLUMN {col} TEXT DEFAULT ''")
        # anything interrupted by a restart goes back in the queue
        for r in c.execute("SELECT id FROM materials WHERE status NOT IN ('ready','error')"):
            work_queue.append((r["id"], True))
        # ready materials from before page-aware search (or missing vectors) just get re-indexed, no AI rerun
        for r in c.execute("SELECT id FROM materials WHERE status='ready' AND (id NOT IN (SELECT material_id FROM "
                           "passages) OR id NOT IN (SELECT material_id FROM vectors))"):
            work_queue.append((r["id"], False))


def material(mid):
    with db() as c:
        r = c.execute("SELECT * FROM materials WHERE id=?", (mid,)).fetchone()
    if not r:
        return None
    d = dict(r)
    d["topics"] = json.loads(d["topics"] or "[]")
    d["notes"] = json.loads(d["notes"] or "[]")
    return d


def update(mid, **fields):
    for k in ("topics", "notes"):
        if k in fields:
            fields[k] = json.dumps(fields[k])
    sets = ", ".join(f"{k}=?" for k in fields)
    with db() as c:
        c.execute(f"UPDATE materials SET {sets} WHERE id=?", (*fields.values(), mid))


# ---------------- text extraction ----------------

def extract_pages(path: Path):
    """The material as [(page number or None, location label, text)], so answers can cite pages and slides."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        out = subprocess.run(["pdftotext", "-enc", "UTF-8", str(path), "-"], capture_output=True, timeout=300)
        pages = out.stdout.decode("utf-8", "replace").split("\x0c")    # pdftotext ends each page with a form feed
        return [(i, f"page {i}", t) for i, t in enumerate(pages, 1) if t.strip()]
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            if texts:
                out.append((i, f"slide {i}", "\n".join(texts)))
        return out
    return [(None, "", extract_text(path))]


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        out = subprocess.run(["pdftotext", "-enc", "UTF-8", str(path), "-"], capture_output=True, timeout=300)
        return out.stdout.decode("utf-8", "replace")
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            if texts:
                out.append(f"Slide {i}:\n" + "\n".join(texts))
        return "\n\n".join(out)
    if ext == ".ipynb":
        nb = json.loads(path.read_text(errors="replace"))
        out = []
        for cell in nb.get("cells", []):
            src = "".join(cell.get("source", []))
            if src.strip():
                out.append(src if cell.get("cell_type") == "markdown" else f"```python\n{src}\n```")
        return "\n\n".join(out)
    return path.read_text(errors="replace")


def clean(text):
    text = text.replace("\x0c", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def split(text, size):
    """Split on paragraph boundaries into pieces of roughly `size` characters."""
    pieces, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        while len(para) > size:                       # very long paragraph: hard cut
            if cur.strip():
                pieces.append(cur.strip())
            pieces.append(para[:size])
            cur, para = "", para[size:]
        if len(cur) + len(para) > size and cur:
            pieces.append(cur.strip())
            cur = ""
        cur += para + "\n\n"
    if cur.strip():
        pieces.append(cur.strip())
    return pieces


# ---------------- local AI (Ollama) ----------------

def ensure_ollama():
    try:
        urllib.request.urlopen(OLLAMA + "/api/version", timeout=2)
        return True
    except Exception:
        subprocess.Popen(["ollama", "serve"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                         start_new_session=True)
        for _ in range(20):
            time.sleep(0.5)
            try:
                urllib.request.urlopen(OLLAMA + "/api/version", timeout=2)
                return True
            except Exception:
                pass
    return False


def strip_think(s):
    # this Ollama build ignores think=false for qwen3, and the opening <think> tag can be missing
    return s.split("</think>")[-1].replace("<think>", "").strip()


def llm(messages, fmt=None, stream=False, temperature=0.3, num_ctx=8192, tries=3):
    """Ask the local model. Non-streaming calls retry, since Ollama can drop a request when RAM is tight."""
    for attempt in range(1, tries + 1):
        try:
            return _llm(messages, fmt, stream, temperature, num_ctx)
        except (urllib.error.URLError, ConnectionError, TimeoutError, json.JSONDecodeError) as e:
            if stream or attempt == tries:
                raise
            print(f"Ollama request failed ({e}); retrying {attempt}/{tries - 1}", flush=True)
            time.sleep(5 * attempt)
            ensure_ollama()


def claude_key():
    return settings().get("anthropic_api_key") or os.environ.get("ANTHROPIC_API_KEY", "")


def strict_schema(schema):
    """Structured outputs want every object closed and no length limits."""
    if isinstance(schema, dict):
        out = {k: strict_schema(v) for k, v in schema.items() if k not in ("maxItems", "minItems")}
        if out.get("type") == "object":
            out["additionalProperties"] = False
        return out
    return schema


def _claude(messages, fmt, stream):
    import anthropic
    client = anthropic.Anthropic(api_key=claude_key() or None, max_retries=3)
    system = "\n\n".join(m["content"] for m in messages if m["role"] == "system")
    convo = [m for m in messages if m["role"] != "system"]
    config = {"effort": "low" if fmt else "medium"}
    if fmt:
        config["format"] = {"type": "json_schema", "schema": strict_schema(fmt)}
    kwargs = dict(model=CLAUDE_MODEL, max_tokens=16000, system=system, messages=convo, output_config=config,
                  betas=["server-side-fallback-2026-07-01"], fallbacks="default")
    if not stream:
        with client.beta.messages.stream(**kwargs) as s:
            msg = s.get_final_message()
        if msg.stop_reason == "refusal":
            raise RuntimeError("Claude declined this request.")
        return "".join(b.text for b in msg.content if b.type == "text")

    def gen():
        with client.beta.messages.stream(**kwargs) as s:
            yield from s.text_stream
            if s.get_final_message().stop_reason == "refusal":
                yield "\n\n⚠️ Claude declined to answer this."
    return gen()


def _llm(messages, fmt, stream, temperature, num_ctx):
    if settings()["model"] == CLAUDE_MODEL:
        return _claude(messages, fmt, stream)
    if messages and messages[0]["role"] == "system":      # qwen3's switch to skip its reasoning step
        messages = [{**messages[0], "content": messages[0]["content"] + "\n/no_think"}] + messages[1:]
    body = {"model": settings()["model"], "messages": messages, "stream": stream, "think": False,
            "keep_alive": "30m", "options": {"num_ctx": num_ctx, "temperature": temperature,
                                             "num_thread": max(2, (os.cpu_count() or 4) // 2)}}
    if fmt:
        body["format"] = fmt
    req = urllib.request.Request(OLLAMA + "/api/chat", data=json.dumps(body).encode(),
                                 headers={"Content-Type": "application/json"})
    resp = urllib.request.urlopen(req, timeout=3600)
    if not stream:
        return strip_think(json.loads(resp.read())["message"]["content"])

    def gen():
        # hold text back until we know whether it's reasoning (ends with </think>) or the answer
        buf, released = "", False
        for line in resp:
            if not line.strip():
                continue
            d = json.loads(line)
            piece = d.get("message", {}).get("content", "")
            if released:
                yield piece
            else:
                buf += piece
                if "</think>" in buf:
                    buf, released = buf.split("</think>", 1)[1].lstrip(), True
                elif len(buf) > 4000:          # long and no </think>: the model skipped reasoning
                    buf, released = buf.replace("<think>", ""), True
                if released and buf:
                    yield buf
            if d.get("done"):
                break
        if not released:
            yield strip_think(buf)
    return gen()


embed_ok = {"v": None}      # None = not checked yet, False = model unavailable (keyword search only)


def embed(texts, query=False):
    """Meaning vectors from the local embedding model, normalized so a dot product is cosine similarity."""
    import numpy as np
    if embed_ok["v"] is None:
        try:
            with urllib.request.urlopen(OLLAMA + "/api/tags", timeout=5) as r:
                have = any(m["name"].split(":")[0] == EMBED_MODEL for m in json.loads(r.read())["models"])
            if not have:      # one-time free download (~270 MB)
                req = urllib.request.Request(OLLAMA + "/api/pull", headers={"Content-Type": "application/json"},
                                             data=json.dumps({"model": EMBED_MODEL, "stream": False}).encode())
                urllib.request.urlopen(req, timeout=1800).read()
            embed_ok["v"] = True
        except Exception as e:
            print(f"Embedding model unavailable ({e}); using keyword search only", flush=True)
            embed_ok["v"] = False
    if not embed_ok["v"]:
        return None
    prefix = "search_query: " if query else "search_document: "
    out = []
    for i in range(0, len(texts), 32):
        body = {"model": EMBED_MODEL, "input": [prefix + t[:2000] for t in texts[i:i + 32]], "keep_alive": "30m"}
        req = urllib.request.Request(OLLAMA + "/api/embed", data=json.dumps(body).encode(),
                                     headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=600) as r:
            out += json.loads(r.read())["embeddings"]
    v = np.array(out, dtype=np.float32)
    return v / np.maximum(np.linalg.norm(v, axis=1, keepdims=True), 1e-9)


def wait_for_chats():
    """Background summarizing pauses while you're talking to the tutor, so replies stay fast."""
    while chat_count["n"] > 0:
        time.sleep(0.5)


# ---------------- processing pipeline ----------------

def classify(m, text):
    with db() as c:
        subjects = [r[0] for r in c.execute(
            "SELECT DISTINCT subject FROM materials WHERE subject IS NOT NULL AND subject!='' AND id!=?", (m["id"],))]
    mid = len(text) // 2
    sample = text[:2500] + ("\n...\n" + text[mid:mid + 1000] if len(text) > 4000 else "")
    schema = {"type": "object", "properties": {
        "title": {"type": "string"},
        "subject": {"type": "string"},
        "kind": {"type": "string", "enum": KINDS},
        "topics": {"type": "array", "items": {"type": "string"}, "maxItems": 8},
        "unit": {"type": "string"},
        "semester": {"type": "string"}},
        "required": ["title", "subject", "kind", "topics", "unit", "semester"]}
    prompt = (f"Existing subjects in the library: {', '.join(subjects) if subjects else 'none yet'}\n"
              f"File name: {m['filename']}\n\nExcerpt of the material:\n{sample}\n\n"
              "Sort this study material. Return JSON with:\n"
              "- title: a short clear title for it\n"
              "- subject: the academic subject or course it belongs to (e.g. 'Machine Learning', "
              "'Physics', 'Data Structures'). If one of the existing subjects fits, reuse that name EXACTLY.\n"
              f"- kind: one of {', '.join(KINDS)}\n"
              "- topics: 3 to 8 main topics it covers, a few words each\n"
              "- unit: the unit, module or chapter, ONLY if the file name or text states it (like 'Unit 2' or "
              "'Chapter 5'), otherwise an empty string\n"
              "- semester: ONLY if stated (like 'Semester 4'), otherwise an empty string")
    out = json.loads(llm([{"role": "system", "content": "You organize a student's study library."},
                          {"role": "user", "content": prompt}], fmt=schema, temperature=0))
    subject = out["subject"].strip()
    for s in subjects:                                   # merge near-duplicates like "machine learning"
        if s.lower() == subject.lower():
            subject = s
    return {"title": out["title"].strip()[:120] or m["filename"], "subject": subject[:60] or "General",
            "kind": out["kind"] if out["kind"] in KINDS else "Other",
            "topics": [t.strip() for t in out["topics"] if t.strip()][:8],
            "unit": tidy_label(out.get("unit"), ("unit", "module", "chapter", "part", "lecture", "week")),
            "semester": tidy_label(out.get("semester"), ("sem",))}


def tidy_label(s, words):
    """Keep 'Unit 3' style labels; small models sometimes fill these with guesses like 'N/A' or a topic."""
    s = (s or "").strip()[:30]
    return s if s and any(w in s.lower() for w in words) and re.search(r"\d|[ivx]+\b", s.lower()) else ""


def summarize(mid, text):
    size = min(PART_MAX, max(3500, math.ceil(len(text) / MAX_PARTS)))
    parts = split(text, size)
    if len(parts) > MAX_PARTS:                           # extremely long: read evenly spaced parts
        step = len(parts) / MAX_PARTS
        parts = [parts[int(i * step)] for i in range(MAX_PARTS)]
    sys_msg = {"role": "system", "content": "You are an expert teacher writing clear, accurate study notes. "
               "Only use what is in the given text; never invent facts."}
    notes = []
    if len(parts) > 1:
        for i, part in enumerate(parts, 1):
            wait_for_chats()
            update(mid, progress=f"Reading part {i} of {len(parts)}…")
            notes.append(llm([sys_msg, {"role": "user", "content":
                f"Section {i} of {len(parts)} of the material:\n\n{part}\n\n"
                "Write compact study notes for this section as bullet points: key concepts, definitions, "
                "formulas, important facts and examples. At most 150 words."}]))
            update(mid, notes=notes)
        source = "\n\n".join(f"Notes on section {i}:\n{n}" for i, n in enumerate(notes, 1))
    else:
        source = parts[0] if parts else ""
    wait_for_chats()
    update(mid, progress="Writing the summary…")
    summary = llm([sys_msg, {"role": "user", "content": f"{source}\n\n"
        "Using the material above, write a study summary in Markdown with these sections:\n"
        "## Overview\n2-3 sentences on what this material is about.\n"
        "## Key concepts\nBullets, each with a one-line plain-language explanation.\n"
        "## Definitions & formulas\nOnly ones that appear in the material (skip the section if none).\n"
        "## Focus for exams\nThe 3-5 most important things to master."}])
    return summary, notes


def build_index(mid, m):
    """Extract the text page by page, split it into passages, and embed them for meaning-based search."""
    pages = [(p, loc, clean(t)) for p, loc, t in extract_pages(LIB / mid / m["stored"])]
    text = "\n\n".join(t for _, _, t in pages if t)
    if len(text) < 40:
        raise ValueError("No readable text found. If it's a scanned PDF or image, export it with text "
                         "(OCR) and upload again.")
    (LIB / mid / "text.txt").write_text(text)
    rows = [(p, loc, piece) for p, loc, t in pages for piece in split(t, CHUNK) if len(piece) > 20]
    with db() as c:
        c.execute("DELETE FROM passages WHERE material_id=?", (mid,))
        c.execute("DELETE FROM vectors WHERE material_id=?", (mid,))
        c.executemany("INSERT INTO passages(material_id, idx, page, loc, text) VALUES (?,?,?,?,?)",
                      [(mid, i, p or "", loc, t) for i, (p, loc, t) in enumerate(rows)])
    vec_cache.pop(mid, None)
    if ensure_ollama():
        wait_for_chats()
        vecs = embed([t for _, _, t in rows])
        if vecs is not None:
            with db() as c:
                c.executemany("INSERT INTO vectors(material_id, idx, vec) VALUES (?,?,?)",
                              [(mid, i, v.tobytes()) for i, v in enumerate(vecs)])
    return text


def process(mid, full=True):
    """full=False only rebuilds the search index (used for materials added before page citations)."""
    m = material(mid)
    if not m:
        return
    try:
        if not full:
            update(mid, progress="Indexing for search…")
            build_index(mid, m)
            update(mid, progress="")
            return
        update(mid, status="processing", progress="Extracting text and indexing for search…", error="")
        text = build_index(mid, m)
        update(mid, chars=len(text))
        if settings()["model"] != CLAUDE_MODEL and not ensure_ollama():
            raise RuntimeError("Ollama isn't running and couldn't be started. Run: ollama serve")
        wait_for_chats()
        update(mid, progress="Sorting into a subject…")
        info = classify(m, text)
        update(mid, **info)
        summary, notes = summarize(mid, text)
        update(mid, summary=summary, notes=notes, status="ready", progress="")
    except Exception as e:
        update(mid, status="error" if full else m["status"], progress="", error=str(e)[:500])


def worker():
    while True:
        with work_cv:
            while not work_queue:
                work_cv.wait()
            mid, full = work_queue.pop(0)
        process(mid, full)


def enqueue(mid, full=True):
    with work_cv:
        work_queue[:] = [j for j in work_queue if j[0] != mid]
        work_queue.append((mid, full))
        work_cv.notify()


# ---------------- tutor ----------------

STOP = set("""a an the and or but if then of to in on at for with by from is are was were be been this that these
those it its as what which who whom how why when where do does did can could should would will shall may might
me my i you your we our they them their he she his her about explain tell teach give show please more less
some any all also just like into over under than so very not no yes there here have has had""".split())


vec_cache = {}      # material id -> (passage indexes, matrix of normalized vectors)


def material_vectors(mid):
    import numpy as np
    if mid not in vec_cache:
        with db() as c:
            rows = c.execute("SELECT idx, vec FROM vectors WHERE material_id=? ORDER BY idx", (mid,)).fetchall()
        vec_cache[mid] = ([r["idx"] for r in rows],
                          np.stack([np.frombuffer(r["vec"], dtype=np.float32) for r in rows]) if rows else None)
    return vec_cache[mid]


def retrieve(scope, query, k=6):
    """Hybrid search over the student's passages: meaning (embeddings) + exact words (BM25), merged by rank.
    Returns (passages, found) where found says whether anything is clearly about the question."""
    words = [w for w in re.findall(r"[A-Za-z0-9]+", query.lower()) if len(w) > 2 and w not in STOP]
    with db() as c:
        if scope == "all":
            mids = [r[0] for r in c.execute("SELECT id FROM materials")]
        elif scope.startswith("subject:"):
            mids = [r[0] for r in c.execute("SELECT id FROM materials WHERE subject=?", (scope[8:],))]
        else:
            mids = [scope]
        if not mids:
            return [], False
        marks = ",".join("?" * len(mids))
        kw = []
        if words:
            q = " OR ".join(f'"{w}"' for w in dict.fromkeys(words))
            kw = [(r[0], int(r[1])) for r in c.execute(
                f"SELECT material_id, idx FROM passages WHERE passages MATCH ? AND material_id IN ({marks}) "
                f"ORDER BY bm25(passages) LIMIT 20", (q, *mids))]
    sims, sem = {}, []
    qv = embed([query], query=True) if any(material_vectors(m)[1] is not None for m in mids) else None
    if qv is not None:
        for m in mids:
            idxs, mat = material_vectors(m)
            if mat is not None:
                for i, s in zip(idxs, (mat @ qv[0]).tolist()):
                    sims[(m, i)] = s
        sem = sorted(sims, key=sims.get, reverse=True)[:20]
    score = {}
    for ranked in (sem, kw):                              # reciprocal rank fusion
        for r, key in enumerate(ranked):
            score[key] = score.get(key, 0) + 1 / (60 + r)
    best = sorted(score, key=score.get, reverse=True)[:k]
    found = (max(sims.values(), default=0) >= FOUND_SIM) if sims else bool(kw)
    with db() as c:
        if not best:                                      # nothing matched: start of the material(s)
            best = [(r[0], int(r[1])) for r in c.execute(
                f"SELECT material_id, idx FROM passages WHERE material_id IN ({marks}) "
                f"ORDER BY CAST(idx AS INTEGER) LIMIT ?", (*mids, k))]
        meta = {r["id"]: dict(r) for r in c.execute(f"SELECT id, title, filename FROM materials WHERE id IN ({marks})", mids)}
        out = []
        for m, i in best:
            r = c.execute("SELECT page, loc, text FROM passages WHERE material_id=? AND idx=?", (m, i)).fetchone()
            if r:
                out.append({"material_id": m, "idx": i, "title": meta[m]["title"] or meta[m]["filename"],
                            "filename": meta[m]["filename"], "page": r["page"] or None, "loc": r["loc"],
                            "text": r["text"], "sim": round(sims.get((m, i), 0), 3)})
    return out, found


def scope_context(scope):
    with db() as c:
        if scope == "all":
            rows = c.execute("SELECT * FROM materials WHERE status='ready' ORDER BY subject").fetchall()
            label = "the student's whole library"
        elif scope.startswith("subject:"):
            rows = c.execute("SELECT * FROM materials WHERE status='ready' AND subject=?", (scope[8:],)).fetchall()
            label = f"the subject '{scope[8:]}'"
        else:
            rows = c.execute("SELECT * FROM materials WHERE id=?", (scope,)).fetchall()
            label = None
    if label is None and rows:
        m = rows[0]
        return (f"Material: {m['title'] or m['filename']} ({m['kind']}, subject: {m['subject']})\n"
                f"Topics: {', '.join(json.loads(m['topics'] or '[]'))}\n\nSummary:\n{m['summary'][:3000]}")
    lines = [f"You are tutoring on {label}. Materials:"]
    budget = 3500 // max(1, len(rows))
    for m in rows[:25]:
        lines.append(f"- {m['title']} ({m['kind']}, {m['subject']}): topics {', '.join(json.loads(m['topics'] or '[]'))}. "
                     f"{re.sub(r'\s+', ' ', m['summary'])[:budget]}")
    return "\n".join(lines)


TUTOR = """You are Study Buddy, a patient, encouraging tutor. You teach the student from THEIR study material.
- Base your teaching on the material context and excerpts below. If you add something that isn't in the
  material, say "(beyond your notes)".
- Explain simply first, then go deeper. Use small examples, analogies, and step-by-step working.
- When teaching a topic, end with ONE short question to check understanding, and wait for the answer.
- When the student answers, tell them kindly whether it's right and why, then continue.
- Use Markdown (headings, bullets, **bold**, code blocks). Keep each reply focused, not a wall of text.
- The excerpts are numbered. When you use one, cite it right after the sentence like [1] or [2]. Never cite a
  number that isn't in the list.
- Follow the student's requests about style: beginner level, Hinglish, real-world examples, comparisons,
  step-by-step working, or "give me a hint, not the solution"."""

NOT_FOUND = """NOTE: the search found NO passage in the student's material that is clearly about this message.
If the student is asking about a concept or fact, start your reply with exactly:
"I couldn't find this in your uploaded material." Then you may give a short general explanation, and mark it
"(beyond your notes)". If the message is a follow-up in the conversation (an answer, "next", "simpler"…), just
continue normally."""

SOURCE_MARK = "\x1e"     # separates the streamed answer from the JSON list of sources at the end


def cited_sources(answer, excerpts, found):
    """The excerpts the answer cites as [n]; if it cited none but the material clearly matched, the top two."""
    nums = [int(n) for n in re.findall(r"\[(\d{1,2})\]", answer)]
    picked = [excerpts[n - 1] for n in dict.fromkeys(nums) if 1 <= n <= len(excerpts)]
    if not picked and found and "couldn't find this in your uploaded material" not in answer.lower():
        picked = excerpts[:2]
    return [{"n": excerpts.index(e) + 1, "material_id": e["material_id"], "idx": e["idx"], "title": e["title"],
             "filename": e["filename"], "page": e["page"], "loc": e["loc"]} for e in picked]


@app.post("/api/chat")
def chat():
    body = request.get_json(force=True)
    scope, msg = body.get("scope", "all"), (body.get("message") or "").strip()
    if not msg:
        abort(400)
    if settings()["model"] != CLAUDE_MODEL and not ensure_ollama():
        return Response("⚠️ Ollama isn't running. Start it with `ollama serve`.", mimetype="text/plain")
    with db() as c:
        history = [dict(r) for r in c.execute(
            "SELECT role, content FROM messages WHERE scope=? ORDER BY id DESC LIMIT 8", (scope,))][::-1]
        c.execute("INSERT INTO messages(scope, role, content, created) VALUES (?,?,?,?)",
                  (scope, "user", msg, time.time()))
    # short follow-ups ("next", "simpler", a quiz answer) search with the last reply for context
    last = history[-1]["content"][:600] if history and len(msg) < 80 else ""
    excerpts, found = retrieve(scope, f"{msg} {last}".strip())
    ctx = scope_context(scope) + "\n\nNumbered excerpts from the material:\n" + "\n\n".join(
        f"[{n}] {e['title']}{' — ' + e['loc'] if e['loc'] else ''}\n{e['text']}" for n, e in enumerate(excerpts, 1))
    system = TUTOR + "\n\n" + ctx + ("" if found else "\n\n" + NOT_FOUND)
    messages = [{"role": "system", "content": system}] + history + [{"role": "user", "content": msg}]

    @stream_with_context
    def gen():
        with chat_lock:
            chat_count["n"] += 1
        answer = ""
        try:
            for piece in llm(messages, stream=True, temperature=0.5):
                answer += piece
                yield piece
        except Exception as e:
            err = f"\n\n⚠️ The tutor stopped: {e}"
            answer += err
            yield err
        finally:
            with chat_lock:
                chat_count["n"] -= 1
            sources = cited_sources(answer, excerpts, found)
            if answer.strip():           # saved even if the browser tab closed mid-answer
                with db() as c:
                    c.execute("INSERT INTO messages(scope, role, content, sources, created) VALUES (?,?,?,?,?)",
                              (scope, "assistant", answer.strip(), json.dumps(sources), time.time()))
        if sources:
            yield SOURCE_MARK + json.dumps(sources)

    return Response(gen(), mimetype="text/plain", headers={"X-Accel-Buffering": "no", "Cache-Control": "no-cache"})


# ---------------- API ----------------

PUBLIC = ("id", "filename", "title", "subject", "kind", "unit", "semester", "topics", "status", "progress", "error",
          "chars", "created")


@app.get("/api/library")
def library():
    with db() as c:
        rows = c.execute("SELECT * FROM materials ORDER BY created DESC").fetchall()
    items = []
    for r in rows:
        d = {k: r[k] for k in PUBLIC}
        d["topics"] = json.loads(d["topics"] or "[]")
        items.append(d)
    return jsonify(items=items, kinds=KINDS, queue=len(work_queue))


@app.post("/api/upload")
def upload():
    added, skipped = [], []
    for f in request.files.getlist("files"):
        name = secure_filename(f.filename or "") or "file.txt"
        ext = Path(name).suffix.lower()
        if ext not in EXTS:
            skipped.append(f"{f.filename}: unsupported type")
            continue
        data = f.read()
        sha = hashlib.sha256(data).hexdigest()
        with db() as c:
            dup = c.execute("SELECT title, filename FROM materials WHERE sha=?", (sha,)).fetchone()
        if dup:
            skipped.append(f"{f.filename}: already in your library")
            continue
        mid = uuid.uuid4().hex[:12]
        (LIB / mid).mkdir(parents=True)
        (LIB / mid / name).write_bytes(data)
        with db() as c:
            c.execute("INSERT INTO materials(id, filename, stored, sha, title, subject, kind, status, progress, created) "
                      "VALUES (?,?,?,?,?,?,?,?,?,?)",
                      (mid, f.filename, name, sha, Path(f.filename).stem, "", "", "queued", "Waiting…", time.time()))
        enqueue(mid)
        added.append(mid)
    return jsonify(added=added, skipped=skipped)


@app.get("/api/material/<mid>")
def get_material(mid):
    m = material(mid) or abort(404)
    m.pop("sha", None)
    return jsonify(m)


@app.patch("/api/material/<mid>")
def edit_material(mid):
    material(mid) or abort(404)
    body = request.get_json(force=True)
    fields = {k: str(body[k]).strip()[:120] for k in ("title", "subject", "kind") if k in body and str(body[k]).strip()}
    fields.update({k: str(body[k]).strip()[:30] for k in ("unit", "semester") if k in body})
    if fields:
        update(mid, **fields)
    return jsonify(ok=True)


@app.post("/api/material/<mid>/reprocess")
def reprocess(mid):
    material(mid) or abort(404)
    update(mid, status="queued", progress="Waiting…", error="")
    enqueue(mid)
    return jsonify(ok=True)


@app.delete("/api/material/<mid>")
def delete_material(mid):
    material(mid) or abort(404)
    with db() as c:
        c.execute("DELETE FROM materials WHERE id=?", (mid,))
        c.execute("DELETE FROM passages WHERE material_id=?", (mid,))
        c.execute("DELETE FROM vectors WHERE material_id=?", (mid,))
        c.execute("DELETE FROM messages WHERE scope=?", (mid,))
    vec_cache.pop(mid, None)
    for p in sorted((LIB / mid).glob("*")):
        p.unlink()
    (LIB / mid).rmdir()
    return jsonify(ok=True)


@app.get("/api/material/<mid>/file")
def material_file(mid):
    m = material(mid) or abort(404)
    return send_from_directory(LIB / mid, m["stored"], download_name=m["filename"])


@app.get("/api/messages")
def messages():
    scope = request.args.get("scope", "all")
    with db() as c:
        rows = c.execute("SELECT role, content, sources FROM messages WHERE scope=? ORDER BY id", (scope,)).fetchall()
    return jsonify([{**dict(r), "sources": json.loads(r["sources"] or "[]")} for r in rows])


@app.delete("/api/messages")
def clear_messages():
    with db() as c:
        c.execute("DELETE FROM messages WHERE scope=?", (request.args.get("scope", "all"),))
    return jsonify(ok=True)


@app.route("/api/settings", methods=["GET", "POST"])
def api_settings():
    if request.method == "POST":
        s = settings()
        body = request.get_json(force=True)
        if "anthropic_api_key" in body:          # kept only in settings.json on this computer
            s["anthropic_api_key"] = str(body["anthropic_api_key"]).strip()
        if body.get("model") in MODELS:
            s["model"] = body["model"]
        SETTINGS_PATH.write_text(json.dumps(s, indent=2))
        os.chmod(SETTINGS_PATH, 0o600)
    try:
        with urllib.request.urlopen(OLLAMA + "/api/tags", timeout=2) as r:
            installed = [m["name"] for m in json.loads(r.read())["models"]]
        up = True
    except Exception:
        installed, up = [], False
    s = settings()
    s.pop("anthropic_api_key", None)
    return jsonify(**s, models=MODELS, installed=installed + ([CLAUDE_MODEL] if claude_key() else []),
                   ollama=up, claude=CLAUDE_MODEL, claude_key=bool(claude_key()))


@app.get("/api/material/<mid>/passage/<int:idx>")
def passage(mid, idx):
    """One cited passage plus its neighbours, for the source viewer."""
    m = material(mid) or abort(404)
    with db() as c:
        rows = c.execute("SELECT idx, page, loc, text FROM passages WHERE material_id=? AND CAST(idx AS INTEGER) "
                         "BETWEEN ? AND ? ORDER BY CAST(idx AS INTEGER)", (mid, idx - 1, idx + 1)).fetchall()
    return jsonify(title=m["title"] or m["filename"], filename=m["filename"], stored=m["stored"],
                   passages=[{**dict(r), "idx": int(r["idx"]), "hit": int(r["idx"]) == idx} for r in rows])


@app.get("/api/search")
def search():
    q = request.args.get("q", "").strip()
    if not q:
        return jsonify(results=[], found=False)
    results, found = retrieve(request.args.get("scope", "all"), q, k=10)
    return jsonify(results=results, found=found)


def scope_label(scope, titles):
    if scope == "all":
        return "Whole library"
    if scope.startswith("subject:"):
        return scope[8:]
    return titles.get(scope, "Deleted material")


@app.get("/api/dashboard")
def dashboard():
    with db() as c:
        titles = {r["id"]: r["title"] or r["filename"] for r in c.execute("SELECT id, title, filename FROM materials")}
        asked = c.execute("SELECT scope, content, created FROM messages WHERE role='user' ORDER BY id DESC").fetchall()
    days = {time.strftime("%Y-%m-%d", time.localtime(r["created"])) for r in asked}
    streak, day = 0, time.time()
    if time.strftime("%Y-%m-%d", time.localtime(day)) not in days:
        day -= 86400                                      # today not studied yet: streak still counts until midnight
    while time.strftime("%Y-%m-%d", time.localtime(day)) in days:
        streak, day = streak + 1, day - 86400
    recent, seen = [], set()
    for r in asked:                                       # latest question per conversation
        if r["scope"] not in seen and len(recent) < 6:
            seen.add(r["scope"])
            recent.append({"scope": r["scope"], "label": scope_label(r["scope"], titles),
                           "question": r["content"][:140], "created": r["created"]})
    return jsonify(questions=len(asked), study_days=len(days), streak=streak, recent=recent)


@app.get("/")
def index():
    return send_from_directory(WEB, "index.html")


@app.get("/<path:p>")
def static_files(p):
    return send_from_directory(WEB, p)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--port", type=int, default=8001)
    args = ap.parse_args()
    init_db()
    threading.Thread(target=worker, daemon=True).start()
    threading.Thread(target=ensure_ollama, daemon=True).start()
    print(f"Study Buddy on http://127.0.0.1:{args.port}")
    app.run(host="127.0.0.1", port=args.port, threaded=True)
